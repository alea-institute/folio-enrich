"""WordIngestor must never treat request content as a server filesystem path.

python-docx's ``Document(x)`` opens ``x`` as a path when it is a str. The old
fallback passed undecodable request content straight to it, so a body like
"/etc/passwd" read files off the server.
"""
from __future__ import annotations

import base64
import io
from unittest.mock import patch

import pytest

from app.models.document import DocumentInput
from app.services.ingestion.word_ingestor import WordIngestor

_HOSTILE = ["/etc/passwd", "../x.docx", "../../etc/shadow", "C:\\\\Windows\\\\win.ini", "x.docx"]


@pytest.mark.parametrize("content", _HOSTILE)
def test_path_like_content_never_reaches_document_as_path(content):
    with patch("docx.Document", side_effect=ValueError("not a docx")) as mock_doc:
        with pytest.raises(ValueError, match="Word document"):
            WordIngestor().ingest(DocumentInput(content=content))
    for call in mock_doc.call_args_list:
        (arg,) = call.args
        assert not isinstance(arg, str), f"Document() got a path: {arg!r}"
        assert isinstance(arg, io.BytesIO)


@pytest.mark.parametrize("content", _HOSTILE)
def test_path_like_content_raises_clear_error_with_real_docx(content):
    with pytest.raises(ValueError, match="base64-encoded .docx"):
        WordIngestor().ingest(DocumentInput(content=content))


def test_valid_docx_still_ingests():
    from docx import Document

    d = Document()
    d.add_paragraph("The tenant shall pay rent monthly.")
    buf = io.BytesIO()
    d.save(buf)
    b64 = base64.b64encode(buf.getvalue()).decode()
    assert "tenant shall pay rent" in WordIngestor().ingest(DocumentInput(content=b64))
